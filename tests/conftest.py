import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "scripts"))

import playbook as pb  # noqa: E402

# A tiny fictional NDA: one numbered section per playbook topic. Party names are invented.
SECTIONS = [
    "Each party may disclose Confidential Information to the other, and each party's obligations as a Receiving Party apply equally.",
    "\"Confidential Information\" means non-public information disclosed by a Disclosing Party, whether written or oral, that is marked or would reasonably be understood to be confidential.",
    "Confidential Information excludes information that is public through no fault of the Receiving Party, already known to it, independently developed, or received from a third party without a duty of confidence.",
    "This Agreement lasts two years. Confidentiality obligations survive for three years after it ends, and for trade secrets as long as they remain trade secrets.",
    "The Receiving Party will use Confidential Information only to evaluate a possible services relationship between the parties.",
    "The Receiving Party may share Confidential Information with its Representatives who need to know it and are bound by duties at least as protective as this Agreement.",
    "If legally compelled to disclose, the Receiving Party will give prompt notice where lawful and reasonably cooperate in seeking a protective order.",
    "On request, the Receiving Party will return or destroy Confidential Information, except copies kept under law or in routine backups, which remain confidential.",
    "Nothing in this Agreement grants either party a right to use Confidential Information retained in unaided memory.",
    "This Agreement does not restrict either party from hiring or competing.",
    "Each party may seek injunctive relief for a breach, in addition to other remedies.",
    "This Agreement is governed by the laws of the State of Delaware, and the parties submit to the courts in Wilmington, Delaware.",
    "Neither party may assign this Agreement without the other party's consent, except to a successor of its entire business.",
    "No license is granted, all information is provided as is, and neither party is obliged to proceed with any transaction.",
]


def make_docx(path: Path, sections=SECTIONS, title="MUTUAL NONDISCLOSURE AGREEMENT") -> Path:
    import docx

    d = docx.Document()
    d.add_paragraph(title)
    d.add_paragraph("This Agreement is between Fernhollow Analytics LLC and Quillbrook Partners Inc.")
    for i, text in enumerate(sections, 1):
        d.add_paragraph(f"{i}. {text}")
    d.save(str(path))
    return path


@pytest.fixture
def topics():
    return pb.parse()


@pytest.fixture
def nda(tmp_path):
    return make_docx(tmp_path / "fixture_nda.docx")


@pytest.fixture
def valid_review(nda, topics):
    findings = []
    for i, text in enumerate(SECTIONS, 1):
        findings.append({
            "topic_id": i,
            "topic_name": topics[i].name,
            "status": "acceptable",
            "clause_found": True,
            "quote": text.split(". ")[0].rstrip(".") if i != 2 else text,
            "section_ref": str(i),
            "rationale": "Meets the preferred position.",
            "proposed_language": "",
            "confidence": "high",
        })
    return {"source_file": str(nda), "playbook_version": pb.version(), "model": "test", "findings": findings}
