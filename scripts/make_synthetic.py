"""Generate the synthetic test set: 15 mutual NDAs (.docx), the seeded-deviation manifest, a draft of the
gold labels, and the manual timing sheet.

    python scripts/make_synthetic.py

Everything here is invented: party names, cities, projects, and dates. Each NDA is assembled from clause
templates for the 14 playbook topics plus ordinary boilerplate, in one of three drafting styles:

- "section": Definitions section with (a), (b) items, "Section N. Title." headings, curly quotes,
  "Disclosing Party" / "Receiving Party".
- "dotted":  UPPERCASE headings, N.1 / N.2 subsections, assignment and governing law inside a
  Miscellaneous section, "Discloser" / "Recipient".
- "runin":   no definitions section, run-in headings ("4. Use. The ..."), plain-English "will".

Planted deviations replace or remove a topic's clause. The manifest records what was planted and the
section it sits in. The gold labels are derived from the manifest and marked unreviewed; the repo owner
reviews and corrects them by hand. Once gold_labels.json is marked reviewed, this script will not
overwrite it.
"""

from __future__ import annotations

import csv
import json
import sys
from dataclasses import dataclass, field
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import playbook as pb  # noqa: E402

DATA = pb.ROOT / "data"
NDA_DIR = DATA / "synthetic_ndas"

# ---------------------------------------------------------------------------------------------
# Clause templates. Placeholders: {CI} {DP} {RP} {shall} {Company} {Counterparty} {project}
# {term} {survival} {law} {venue}. Every template is one paragraph.
# ---------------------------------------------------------------------------------------------

ACCEPTABLE = {
    "T1": ('Each party may disclose {CI} to the other party in connection with the Purpose. A party disclosing '
           '{CI} is the "{DP}", and a party receiving it is the "{RP}". Each party, when acting as {RP}, {shall} '
           'hold the other party\'s {CI} in strict confidence, and the obligations in this Agreement apply equally '
           'to both parties.'),
    "T2": ('"{CI}" means all non-public business, technical, and financial information disclosed by or on behalf '
           'of a {DP} to the {RP}, in any form, whether written, electronic, oral, or visual, that is marked or '
           'identified as confidential or that a reasonable person would understand to be confidential given the '
           'nature of the information and the circumstances of disclosure. Oral and visual disclosures are '
           'covered whether or not they are later summarized in writing.'),
    "purpose": ('"Purpose" means the evaluation, negotiation, and performance of a potential services relationship '
                'between the parties concerning {project}.'),
    "T3": ('The obligations in this Agreement do not apply to information that the {RP} can demonstrate by '
           'written records: (a) is or becomes publicly available through no breach of this Agreement by the '
           '{RP} or its Representatives; (b) was known to the {RP} without restriction before disclosure by the '
           '{DP}; (c) is independently developed by the {RP} without use of or reference to the {DP}\'s {CI}; '
           'or (d) is rightfully received by the {RP} from a third party without a duty of confidentiality.'),
    "T4": ('This Agreement covers {CI} disclosed during the {term} following the Effective Date. Either party '
           'may terminate this Agreement earlier on thirty (30) days\' written notice to the other party. The '
           '{RP}\'s obligations under this Agreement survive for {survival} after expiration or termination, and '
           'obligations with respect to trade secrets survive for as long as the information remains a trade '
           'secret under applicable law.'),
    "T5": 'The {RP} {shall} use the {DP}\'s {CI} solely for the Purpose and for no other purpose.',
    "T5_inline": ('The {RP} {shall} use the {DP}\'s {CI} solely for the purpose of evaluating, negotiating, and '
                  'performing a potential services relationship between the parties concerning {project} (the '
                  '"Purpose") and for no other purpose.'),
    "T6": ('The {RP} may disclose {CI} to its and its affiliates\' employees, officers, directors, contractors, '
           'and professional advisors (including legal, accounting, and financial advisors) who need to know it '
           'for the Purpose and who are bound by confidentiality obligations at least as protective as those in '
           'this Agreement (collectively, "Representatives"). The {RP} {shall} not otherwise disclose {CI} to any '
           'third party without the {DP}\'s prior written consent. The {RP} is responsible for any breach of '
           'this Agreement by its Representatives.'),
    "T7": ('If the {RP} or any of its Representatives is required by law, regulation, subpoena, or court order '
           'to disclose any {CI}, the {RP} {shall}, to the extent legally permitted, give the {DP} prompt written '
           'notice so that the {DP} may seek a protective order or other remedy, reasonably cooperate with those '
           'efforts at the {DP}\'s expense, and disclose only the portion of the {CI} that it is legally required '
           'to disclose.'),
    "T8": ('Within thirty (30) days after the {DP}\'s written request, or after this Agreement expires or is '
           'terminated, the {RP} {shall} return or destroy all of the {DP}\'s {CI} in its possession and, on '
           'request, certify the destruction in writing. The {RP} is not required to delete {CI} held in '
           'automatic electronic backup systems or retained to comply with applicable law or its document '
           'retention policies, and any {CI} so retained remains subject to this Agreement for as long as it is '
           'retained.'),
    "T9": ('Nothing in this Agreement permits either party to use {CI} because it has been retained in the memory '
           'of its personnel. Information retained in memory is protected in the same way as all other {CI}.'),
    "T11": ('Each party acknowledges that unauthorized use or disclosure of the other party\'s {CI} may cause '
            'irreparable harm for which monetary damages would be an inadequate remedy. The non-breaching party '
            'may therefore seek injunctive or other equitable relief, without the necessity of proving actual '
            'damages, in addition to any other remedy available at law or in equity. Each party {shall} bear its '
            'own attorneys\' fees and costs in any dispute under this Agreement.'),
    "T11_prevailing": ('Each party acknowledges that unauthorized use or disclosure of the other party\'s {CI} may '
                       'cause irreparable harm for which monetary damages would be an inadequate remedy. The '
                       'non-breaching party may therefore seek injunctive or other equitable relief, without the '
                       'necessity of proving actual damages, in addition to any other remedy available at law or in '
                       'equity. In any action to enforce this Agreement, the prevailing party is entitled to recover '
                       'its reasonable attorneys\' fees and costs.'),
    "T12": ('This Agreement is governed by the laws of the State of {law}, without regard to its conflict of laws '
            'principles. Each party submits to the exclusive jurisdiction of the state and federal courts located '
            'in {venue} for any action arising out of this Agreement.'),
    "T13": ('Neither party may assign this Agreement without the other party\'s prior written consent, except that '
            'either party may assign this Agreement without consent to an affiliate or to a successor in '
            'connection with a merger, acquisition, or sale of all or substantially all of its assets, on written '
            'notice to the other party. This Agreement binds and benefits the parties and their permitted '
            'successors and assigns.'),
    "T14": ('No license under any patent, copyright, trade secret, or other intellectual property right is granted '
            'or implied by this Agreement or by any disclosure under it. All {CI} is provided "as is", and neither '
            'party makes any warranty, express or implied, as to its accuracy or completeness. Neither party is '
            'obligated to enter into any further agreement or transaction, and either party may end discussions '
            'at any time for any reason.'),
}

BOILERPLATE = [
    ("Notices", 'All notices under this Agreement must be in writing and are effective when delivered by hand, by '
                'nationally recognized courier, or by email with confirmation of receipt, to the address shown '
                'below a party\'s signature or to another address that party designates by notice.'),
    ("Entire Agreement", 'This Agreement is the entire agreement of the parties regarding its subject matter and '
                         'supersedes all prior discussions and understandings regarding that subject matter. It may '
                         'be amended only in a writing signed by both parties.'),
    ("Severability and Waiver", 'If any provision of this Agreement is held unenforceable, the remaining provisions '
                                'remain in effect. No failure or delay in exercising any right under this Agreement '
                                'operates as a waiver of that right.'),
    ("Counterparts", 'This Agreement may be signed in counterparts, including by electronic signature, each of '
                     'which is an original and all of which together are one instrument.'),
    ("Relationship of the Parties", 'The parties are independent contractors. Nothing in this Agreement creates an '
                                    'agency, partnership, or joint venture between them.'),
]

HEADINGS = {
    "defs": "Definitions", "T1": "Mutual Obligations", "T2": "Confidential Information",
    "T3": "Exclusions", "T4": "Term and Survival", "T5": "Use of Confidential Information",
    "T6": "Disclosure to Representatives", "T7": "Compelled Disclosure", "T8": "Return or Destruction",
    "T9": "Information Retained in Memory", "T10": "Non-Solicitation", "T11": "Remedies",
    "T12": "Governing Law and Venue", "T13": "Assignment", "T14": "No License; No Warranty; No Obligation",
    "misc": "Miscellaneous",
}

# Section order for each style. "defs" holds T2 and the Purpose definition in the styles that have one.
ORDER = {
    "section": ["defs", "T1", "T3", "T5", "T6", "T7", "T4", "T8", "T9", "T10", "T14", "T11", "T13", "T12", "boiler"],
    "dotted": ["defs", "T1", "T5", "T6", "T3", "T7", "T8", "T4", "T9", "T10", "T11", "T14", "misc"],
    "runin": ["T1", "T2", "T3", "T5", "T6", "T7", "T8", "T4", "T9", "T10", "T11", "T12", "T13", "T14", "boiler"],
}

TERMS = {
    "section": {"CI": "Confidential Information", "DP": "Disclosing Party", "RP": "Receiving Party", "shall": "shall"},
    "dotted": {"CI": "Confidential Information", "DP": "Discloser", "RP": "Recipient", "shall": "shall"},
    "runin": {"CI": "Confidential Information", "DP": "Disclosing Party", "RP": "Receiving Party", "shall": "will"},
}


@dataclass
class Party:
    full: str
    short: str
    entity: str     # e.g. "a Delaware limited liability company"
    city: str       # e.g. "Chicago, Illinois"

    @property
    def state(self) -> str:
        return self.city.split(", ")[1]


@dataclass
class Deviation:
    item: str            # clause key replaced ("T7", "purpose", "def_parties", ...)
    topic: int
    status: str          # expected gold status
    planted: str         # what was planted, in plain words
    text: str | None     # replacement text; None removes the clause


@dataclass
class NDASpec:
    n: int
    kind: str                         # clean | deviation | hard
    company: Party
    counterparty: Party
    company_first: bool
    project: str
    effective: str
    term: str = "two (2) years"
    survival: str = "three (3) years"
    law: tuple[str, str] = ("Delaware", "Wilmington, Delaware")
    fees: str = "own"                 # own | prevailing
    residuals_clause: bool = False    # include the acceptable "no residuals" clause (topic 9 acceptable)
    title: str = ""
    terms: dict = field(default_factory=dict)       # override defined terms
    extra_defs: list[tuple[str, str]] = field(default_factory=list)   # (item key, text) added to definitions
    deviations: list[Deviation] = field(default_factory=list)
    gold_notes: dict[int, str] = field(default_factory=dict)

    @property
    def file(self) -> str:
        return f"nda_{self.n:02d}.docx"

    @property
    def style(self) -> str:
        return ("section", "dotted", "runin")[(self.n - 1) % 3]


def P(full, short, entity, city):
    return Party(full, short, entity, city)


# ---------------------------------------------------------------------------------------------
# The 15 NDAs. Kinds are spread across file numbers so a file name does not reveal its kind.
# The Company (the party using the playbook) is named first in odd-numbered files.
# ---------------------------------------------------------------------------------------------

def specs() -> list[NDASpec]:
    S = []

    S.append(NDASpec(
        1, "deviation",
        P("Fernhollow Analytics LLC", "Fernhollow", "a Delaware limited liability company", "Chicago, Illinois"),
        P("Quillbrook Partners Inc.", "Quillbrook", "an Ohio corporation", "Columbus, Ohio"),
        True, "data analytics and reporting services", "March 2, 2026", law=("Ohio", "Franklin County, Ohio"),
        deviations=[
            Deviation("T7", 7, "fallback", "Compelled disclosure keeps the 'only the portion required' limit but drops "
                      "the notice and cooperation duties.",
                      "If the {RP} is required by law or court order to disclose any {CI}, it may disclose only that "
                      "portion of the {CI} that it is legally required to disclose."),
            Deviation("T12", 12, "fallback", "Governing law is the Counterparty's home state (Ohio, acceptable), but "
                      "venue is exclusive to the Counterparty's county.",
                      "This Agreement is governed by the laws of the State of Ohio, without regard to its conflict of "
                      "laws principles. Any action arising out of this Agreement shall be brought exclusively in the "
                      "state courts located in Franklin County, Ohio."),
        ]))

    S.append(NDASpec(
        2, "deviation",
        P("Tamsworth Logistics Corp.", "Tamsworth", "a Texas corporation", "Fort Worth, Texas"),
        P("Brightvale Facility Services LLC", "Brightvale", "a Delaware limited liability company", "Denver, Colorado"),
        False, "warehouse facilities maintenance services", "January 15, 2026", law=("Delaware", "Wilmington, Delaware"),
        deviations=[
            Deviation("T4", 4, "escalate", "Obligations for all Confidential Information last in perpetuity, not "
                      "only for trade secrets.",
                      "This Agreement remains in effect until either party terminates it on thirty (30) days' "
                      "written notice to the other party. The {RP}'s obligations with respect to all {CI} continue "
                      "in perpetuity and survive any expiration or termination of this Agreement."),
            Deviation("T8", 8, "fallback", "Return or destruction clause has no carve-out for backups or legal "
                      "retention.",
                      "Within thirty (30) days after the {DP}'s written request, the {RP} shall return or destroy all "
                      "of the {DP}'s {CI}, including all copies, extracts, and summaries, and on request certify the "
                      "destruction in writing."),
            Deviation("T14", 14, "fallback", "No-license and no-obligation statements present; the 'as is' no "
                      "warranty statement is missing.",
                      "No license under any patent, copyright, trade secret, or other intellectual property right is "
                      "granted or implied by this Agreement. Neither party is obligated to enter into any further "
                      "agreement or transaction, and either party may end discussions at any time."),
        ]))

    S.append(NDASpec(
        3, "clean",
        P("Cobaltine Software Inc.", "Cobaltine", "a Delaware corporation", "Austin, Texas"),
        P("Marrowfield Consulting Group LLC", "Marrowfield", "a Georgia limited liability company", "Atlanta, Georgia"),
        True, "software implementation consulting services", "February 9, 2026", term="one (1) year",
        survival="two (2) years", law=("New York", "New York County, New York"), fees="prevailing"))

    S.append(NDASpec(
        4, "deviation",
        P("Oakhaven Data Services Inc.", "Oakhaven", "a Delaware corporation", "Raleigh, North Carolina"),
        P("Pellucid Robotics Corp.", "Pellucid", "a Michigan corporation", "Ann Arbor, Michigan"),
        False, "managed cloud hosting services", "April 6, 2026",
        deviations=[
            Deviation("T10", 10, "escalate", "Mutual 24-month non-solicitation of customers and suppliers.",
                      "During the term of this Agreement and for twenty-four (24) months after it ends, neither party "
                      "shall, directly or indirectly, solicit or accept business from any customer or supplier of the "
                      "other party identified to it through the other party's {CI}."),
            Deviation("T11", 11, "escalate", "Equitable relief is mutual, but only the Counterparty may recover "
                      "attorneys' fees, from the Company.",
                      "Each party acknowledges that unauthorized use or disclosure of the other party's {CI} may cause "
                      "irreparable harm, and the non-breaching party may seek injunctive or other equitable relief in "
                      "addition to any other remedy. In any action to enforce this Agreement, {Counterparty} is "
                      "entitled to recover from {Company} its reasonable attorneys' fees and costs."),
            Deviation("T13", 13, "fallback", "Assignment is barred in every case, including a merger or change of "
                      "control.",
                      "Neither party may assign or transfer this Agreement, whether voluntarily, by operation of law, "
                      "by merger, or by change of control, without the other party's prior written consent. Any "
                      "attempted assignment in violation of this Section is void."),
        ]))

    S.append(NDASpec(
        5, "hard",
        P("Sablewood Capital Advisors LLC", "Sablewood", "a New York limited liability company", "New York, New York"),
        P("Thistlecrest Property Management Inc.", "Thistlecrest", "a Florida corporation", "Tampa, Florida"),
        True, "property management and leasing services", "May 4, 2026", law=("New York", "New York County, New York"),
        residuals_clause=False,
        deviations=[
            Deviation("purpose", 5, "escalate", "Hard case (buried in definitions): the defined Purpose adds each "
                      "party's development and improvement of its own products and services. The use clause itself "
                      "reads 'solely for the Purpose'.",
                      '"Purpose" means the evaluation, negotiation, and performance of a potential services '
                      'relationship between the parties concerning {project}, together with each party\'s development '
                      'and improvement of its own products and services.'),
        ]))

    S.append(NDASpec(
        6, "deviation",
        P("Veridune Energy Solutions LLC", "Veridune", "a Colorado limited liability company", "Boulder, Colorado"),
        P("Wexmoor Staffing Partners Inc.", "Wexmoor", "a Delaware corporation", "Phoenix, Arizona"),
        False, "contract staffing services", "June 1, 2026", term="three (3) years", survival="five (5) years",
        deviations=[
            Deviation("T9", 9, "escalate", "Broad residuals clause: any retained information, intentionally "
                      "memorized or not, usable for any purpose with no royalty.",
                      "Notwithstanding anything to the contrary in this Agreement, either party may use for any "
                      "purpose any information, ideas, concepts, know-how, or techniques retained in the memory of its "
                      "personnel, whether or not intentionally memorized, and no royalty or other payment will be owed "
                      "for that use."),
            Deviation("T13", 13, "absent", "Assignment clause removed (absence is acceptable under the playbook; "
                      "not counted as a planted deviation).", None),
        ]))

    S.append(NDASpec(
        7, "deviation",
        P("Lanternfield Media LLC", "Lanternfield", "an Illinois limited liability company", "Evanston, Illinois"),
        P("Halcyra Biologics Inc.", "Halcyra", "a Delaware corporation", "Cambridge, Massachusetts"),
        True, "marketing and creative services", "July 13, 2026", law=("Delaware", "Wilmington, Delaware"),
        deviations=[
            Deviation("T1", 1, "escalate", "Labeled mutual, but the obligation clause binds only the Company.",
                      'Each party may disclose {CI} to the other party in connection with the Purpose. A party '
                      'disclosing {CI} is the "{DP}", and a party receiving it is the "{RP}". {Company} shall hold all '
                      '{CI} disclosed by {Counterparty} in strict confidence and shall not disclose it except as '
                      'permitted by this Agreement.'),
            Deviation("T2", 2, "escalate", "Oral and visual disclosures are protected only if summarized in writing "
                      "within five days.",
                      '"{CI}" means all non-public business, technical, and financial information disclosed by or on '
                      'behalf of a {DP} to the {RP} that is marked or identified as confidential or that a reasonable '
                      'person would understand to be confidential. Information disclosed orally or visually is {CI} '
                      'only if the {DP} summarizes it in writing, marked as confidential, and delivers the summary to '
                      'the {RP} within five (5) days after disclosure.'),
            Deviation("T7", 7, "escalate", "Compelled disclosure clause removed entirely.", None),
            Deviation("T12", 12, "escalate", "Delaware law, but disputes go to binding arbitration.",
                      "This Agreement is governed by the laws of the State of Delaware. Any dispute arising out of or "
                      "relating to this Agreement shall be resolved exclusively by binding arbitration administered in "
                      "Chicago, Illinois under the commercial arbitration rules then in effect, and judgment on the "
                      "award may be entered in any court having jurisdiction."),
        ]))

    S.append(NDASpec(
        8, "clean",
        P("Juniperline Retail Group Inc.", "Juniperline", "a Minnesota corporation", "Minneapolis, Minnesota"),
        P("Kestrelford Insurance Services LLC", "Kestrelford", "a Delaware limited liability company", "Hartford, Connecticut"),
        False, "employee benefits administration services", "August 3, 2026", residuals_clause=True))

    S.append(NDASpec(
        9, "deviation",
        P("Northbask Manufacturing Co.", "Northbask", "a Wisconsin corporation", "Milwaukee, Wisconsin"),
        P("Glimmerton Labs Inc.", "Glimmerton", "a Delaware corporation", "San Jose, California"),
        True, "quality inspection software services", "March 23, 2026", law=("Wisconsin", "Milwaukee County, Wisconsin"),
        deviations=[
            Deviation("T5_inline", 5, "escalate", "Use clause also allows training and improving machine learning "
                      "models.",
                      'The {RP} will use the {DP}\'s {CI} solely for the purpose of evaluating, negotiating, and '
                      'performing a potential services relationship between the parties concerning {project} (the '
                      '"Purpose"), and may also use {CI} to train, test, and improve its machine learning models, '
                      'provided that those models do not disclose {CI} to third parties.'),
            Deviation("T3", 3, "fallback", "Exclusions list omits the independent development exclusion.",
                      "The obligations in this Agreement do not apply to information that the {RP} can demonstrate "
                      "by written records: (a) is or becomes publicly available through no breach of this Agreement "
                      "by the {RP} or its Representatives; (b) was known to the {RP} without restriction before "
                      "disclosure by the {DP}; or (c) is rightfully received by the {RP} from a third party without a "
                      "duty of confidentiality."),
        ]))

    S.append(NDASpec(
        10, "hard",
        P("Duskwater Transit LLC", "Duskwater", "an Oregon limited liability company", "Portland, Oregon"),
        P("Ironpeak Construction Services LLC", "Ironpeak", "a Delaware limited liability company", "Boise, Idaho"),
        False, "fleet maintenance scheduling services", "September 8, 2026",
        title="MUTUAL CONFIDENTIALITY AGREEMENT",
        terms={"DP": "Disclosing Party", "RP": "Receiving Party"},
        extra_defs=[
            ("def_dp", '"Disclosing Party" means {Counterparty} and its affiliates.'),
            ("def_rp", '"Receiving Party" means {Company}.'),
        ],
        deviations=[
            Deviation("T1", 1, "escalate", "Hard case (defined-term trap): titled mutual and the obligation clause "
                      "reads neutrally, but the definitions fix the Counterparty as the only Disclosing Party and the "
                      "Company as the only Receiving Party, so only the Company's obligations exist.",
                      "The Receiving Party shall hold the Disclosing Party's Confidential Information in strict "
                      "confidence and shall not disclose it except as permitted by this Agreement. The parties intend "
                      "this Agreement to protect the information exchanged in their discussions."),
        ],
        gold_notes={11: "Remedies clause is written 'each party', so it stays acceptable; the one-way effect is "
                        "captured under topic 1."}))

    S.append(NDASpec(
        11, "deviation",
        P("Ashgrove Clinical Partners LLC", "Ashgrove", "a Pennsylvania limited liability company", "Pittsburgh, Pennsylvania"),
        P("Copperlane Payroll Inc.", "Copperlane", "a Delaware corporation", "Nashville, Tennessee"),
        True, "payroll processing services", "October 1, 2026", law=("Pennsylvania", "Allegheny County, Pennsylvania"),
        residuals_clause=True,
        deviations=[
            Deviation("T2", 2, "fallback", "Definition requires marking, with no reasonable-person standard.",
                      '"{CI}" means all non-public business, technical, and financial information disclosed by or on '
                      'behalf of a {DP} to the {RP}, in any form, that is clearly marked "Confidential" at the time of '
                      'disclosure. Information disclosed orally or visually is {CI} if it is identified as '
                      'confidential at the time of disclosure.'),
            Deviation("T6", 6, "fallback", "Representatives limited to employees; affiliates, contractors, and "
                      "professional advisors left out.",
                      "The {RP} may disclose {CI} to its employees who need to know it for the Purpose and who are "
                      "bound by written confidentiality obligations at least as protective as those in this Agreement "
                      '(collectively, "Representatives"). The {RP} shall not otherwise disclose {CI} to any third '
                      "party without the {DP}'s prior written consent. The {RP} is responsible for any breach of this "
                      "Agreement by its Representatives."),
        ]))

    S.append(NDASpec(
        12, "clean",
        P("Rookhaven Hospitality Group LLC", "Rookhaven", "a Nevada limited liability company", "Reno, Nevada"),
        P("Saltmeadow Linen Supply Inc.", "Saltmeadow", "a California corporation", "Sacramento, California"),
        False, "commercial linen and laundry services", "May 18, 2026", term="three (3) years", survival="two (2) years",
        law=("Delaware", "Wilmington, Delaware"), fees="prevailing"))

    S.append(NDASpec(
        13, "deviation",
        P("Mistral Point Engineering Inc.", "Mistral Point", "a Washington corporation", "Seattle, Washington"),
        P("Larkspur Survey Group LLC", "Larkspur", "a Delaware limited liability company", "Spokane, Washington"),
        True, "geotechnical survey services", "June 22, 2026", law=("Washington", "King County, Washington"),
        residuals_clause=True,
        deviations=[
            Deviation("T4", 4, "fallback", "Survival is eighteen months (between one and two years) with no "
                      "trade-secret carve-out.",
                      "This Agreement covers {CI} disclosed during the two (2) years following the Effective Date, "
                      "and either party may terminate it earlier on thirty (30) days' written notice. The {RP}'s "
                      "obligations under this Agreement survive for eighteen (18) months after expiration or "
                      "termination."),
            Deviation("T11", 11, "fallback", "Equitable relief runs in one direction only (the Counterparty's "
                      "information); fees are not shifted.",
                      "The {RP} acknowledges that unauthorized use or disclosure of {Counterparty}'s {CI} may cause "
                      "irreparable harm for which monetary damages would be inadequate, and {Counterparty} may seek "
                      "injunctive or other equitable relief in addition to any other remedy. Each party shall bear its "
                      "own attorneys' fees and costs."),
        ]))

    S.append(NDASpec(
        14, "hard",
        P("Briarstone Health Systems Inc.", "Briarstone", "a Delaware corporation", "Indianapolis, Indiana"),
        P("Cindervale IT Partners LLC", "Cindervale", "an Indiana limited liability company", "Fort Wayne, Indiana"),
        False, "help desk and IT support services", "July 27, 2026", law=("Indiana", "Marion County, Indiana"),
        deviations=[
            Deviation("T4", 4, "escalate", "Hard case (ambiguous): no fixed term and no termination right; obligations "
                      "survive 'for such period as is reasonable', so the duration cannot be determined.",
                      "This Agreement continues until the parties' discussions regarding the Purpose have concluded. "
                      "The {RP}'s obligations under this Agreement survive for such period as is reasonable in light "
                      "of the nature of the {CI} concerned."),
        ]))

    S.append(NDASpec(
        15, "deviation",
        P("Hollisford Water Technologies Inc.", "Hollisford", "a Delaware corporation", "Cleveland, Ohio"),
        P("Emberly Field Services LLC", "Emberly", "a Kentucky limited liability company", "Louisville, Kentucky"),
        True, "pump station monitoring services", "August 24, 2026", law=("Ohio", "Cuyahoga County, Ohio"),
        deviations=[
            Deviation("T8", 8, "escalate", "Destruction within five business days, certified by an officer under "
                      "penalty of perjury.",
                      "Within five (5) business days after the {DP}'s written request, the {RP} will return or destroy "
                      "all of the {DP}'s {CI}, and an officer of the {RP} will certify the destruction in writing under "
                      "penalty of perjury. The {RP} is not required to delete {CI} held in automatic electronic backup "
                      "systems, provided that any retained {CI} remains subject to this Agreement."),
            Deviation("T6", 6, "escalate", "Each Representative must sign a joinder before receiving information.",
                      "The {RP} may disclose {CI} to its and its affiliates' employees, officers, directors, "
                      "contractors, and professional advisors who need to know it for the Purpose (collectively, "
                      '"Representatives"), but only after each Representative has signed a joinder to this Agreement '
                      "in a form approved by the {DP}. The {RP} is responsible for any breach of this Agreement by "
                      "its Representatives."),
            Deviation("T14", 14, "escalate", "Feedback and improvements are assigned to the party that receives them.",
                      'No license under any patent, copyright, or other intellectual property right is granted by this '
                      'Agreement, except that all feedback, suggestions, and improvements that either party provides '
                      'regarding the other party\'s products or services are hereby assigned to the party receiving '
                      'them. All {CI} is provided "as is", without warranty of any kind. Neither party is obligated to '
                      'enter into any further agreement or transaction.'),
        ]))

    return S


# ---------------------------------------------------------------------------------------------
# Assembly
# ---------------------------------------------------------------------------------------------

def _ctx(spec: NDASpec) -> dict:
    c = dict(TERMS[spec.style])
    c.update(spec.terms)
    c.update(Company=spec.company.short, Counterparty=spec.counterparty.short, project=spec.project,
             term=spec.term, survival=spec.survival, law=spec.law[0], venue=spec.law[1])
    return c


def _items(spec: NDASpec) -> dict[str, str | None]:
    """Clause key -> text (None = omitted) after deviations are applied."""
    items = {k: ACCEPTABLE[k] for k in ("T1", "T2", "purpose", "T3", "T4", "T5", "T6", "T7", "T8", "T12", "T13", "T14")}
    items["T5_inline"] = ACCEPTABLE["T5_inline"]
    items["T11"] = ACCEPTABLE["T11_prevailing" if spec.fees == "prevailing" else "T11"]
    items["T9"] = ACCEPTABLE["T9"] if spec.residuals_clause else None
    items["T10"] = None
    for d in spec.deviations:
        items[d.item] = d.text
    return items


def build_sections(spec: NDASpec) -> list[tuple[str, list[tuple[str, str]]]]:
    """[(heading, [(item key, text), ...]), ...] in document order, placeholders filled."""
    ctx, items, style = _ctx(spec), _items(spec), spec.style
    fill = lambda k: items[k].format(**ctx)  # noqa: E731
    sections = []
    for key in ORDER[style]:
        if key == "defs":
            entries = [("T2", fill("T2")), ("purpose", fill("purpose"))]
            entries += [(k, t.format(**ctx)) for k, t in spec.extra_defs]
            sections.append((HEADINGS["defs"], entries))
        elif key == "T5" and style == "runin":
            sections.append((HEADINGS["T5"], [("T5_inline", fill("T5_inline"))]))
        elif key == "misc":
            entries = [(k, fill(k)) for k in ("T13", "T12") if items[k] is not None]
            entries += [("boiler", text) for _, text in BOILERPLATE]
            sections.append((HEADINGS["misc"], entries))
        elif key == "boiler":
            sections += [(h, [("boiler", t)]) for h, t in BOILERPLATE]
        elif items.get(key) is not None:
            sections.append((HEADINGS[key], [(key, fill(key))]))
    return sections


def _smart_quotes(text: str) -> str:
    out = []
    for i, ch in enumerate(text):
        prev = text[i - 1] if i else " "
        if ch == '"':
            out.append("“" if prev in " ([" else "”")
        elif ch == "'":
            out.append("’")
        else:
            out.append(ch)
    return "".join(out)


def render(spec: NDASpec) -> tuple[list[tuple[str, str, str]], dict[str, str]]:
    """Paragraph plan [(kind, ref-or-label, text)] and item key -> section ref."""
    style = spec.style
    paras, refs = [], {}
    sub_labels = "abcdefghijklmnop"
    for n, (heading, entries) in enumerate(build_sections(spec), 1):
        if style == "section":
            paras.append(("heading", f"Section {n}.", f"Section {n}. {heading}."))
            for j, (k, text) in enumerate(entries):
                if len(entries) > 1:
                    label = f"({sub_labels[j]})"
                    paras.append(("body", f"{n}({sub_labels[j]})", f"{label} {text}"))
                    refs.setdefault(k, f"{n}({sub_labels[j]})")
                else:
                    paras.append(("body", str(n), text))
                    refs.setdefault(k, str(n))
        elif style == "dotted":
            paras.append(("heading", f"{n}.", f"{n}. {heading.upper()}"))
            for j, (k, text) in enumerate(entries, 1):
                if len(entries) > 1:
                    paras.append(("body", f"{n}.{j}", f"{n}.{j} {text}"))
                    refs.setdefault(k, f"{n}.{j}")
                else:
                    paras.append(("body", str(n), text))
                    refs.setdefault(k, str(n))
        else:  # runin
            (k, text), = entries
            paras.append(("runin", f"{n}. {heading}.", text))
            refs.setdefault(k, str(n))
    return paras, refs


def preamble(spec: NDASpec) -> tuple[str, list[str]]:
    a, b = (spec.company, spec.counterparty) if spec.company_first else (spec.counterparty, spec.company)
    title = spec.title or ("MUTUAL NONDISCLOSURE AGREEMENT" if spec.style != "runin" else "Mutual Confidentiality Agreement")
    opening = (f'This {"Mutual Nondisclosure Agreement" if "NONDISCLOSURE" in title else "Mutual Confidentiality Agreement"} '
               f'(the "Agreement") is entered into as of {spec.effective} (the "Effective Date") by and between '
               f'{a.full}, {a.entity} with offices in {a.city} ("{a.short}"), and {b.full}, {b.entity} with offices '
               f'in {b.city} ("{b.short}").')
    if spec.style == "section":
        recitals = [opening,
                    f"Background. {a.short} and {b.short} wish to explore a potential services relationship concerning "
                    f"{spec.project}. In the course of their discussions each party expects to disclose confidential "
                    f"information to the other, and each party wishes to protect its information on the terms below.",
                    "The parties therefore agree as follows:"]
    elif spec.style == "dotted":
        recitals = [opening, "RECITALS",
                    f"WHEREAS, the parties are considering a business relationship concerning {spec.project};",
                    "WHEREAS, in connection with those discussions each party may disclose confidential information "
                    "to the other; and",
                    "WHEREAS, the parties wish to set out the terms on which that information will be protected.",
                    "NOW, THEREFORE, in consideration of the mutual promises below, the parties agree as follows:"]
    else:
        recitals = [opening,
                    f"The parties are talking about {spec.project} and will share confidential information with each "
                    f"other while they do. They agree as follows."]
    return title, recitals


def write_docx(spec: NDASpec, path: Path) -> dict[str, str]:
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt

    style = spec.style
    d = docx.Document()
    for s in d.sections:
        s.left_margin = s.right_margin = Inches(1)
        s.top_margin = s.bottom_margin = Inches(1)
    normal = d.styles["Normal"]
    normal.font.name = {"section": "Times New Roman", "dotted": "Calibri", "runin": "Georgia"}[style]
    normal.font.size = Pt({"section": 11.5, "dotted": 11, "runin": 11}[style])
    normal.paragraph_format.space_after = Pt(6)
    q = _smart_quotes if style == "section" else (lambda t: t)

    title, recitals = preamble(spec)
    t = d.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t.add_run(title).bold = True
    for r in recitals:
        p = d.add_paragraph(q(r))
        if r == "RECITALS":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    paras, refs = render(spec)
    for kind, label, text in paras:
        if kind == "heading":
            p = d.add_paragraph()
            p.add_run(q(text)).bold = True
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.keep_with_next = True
        elif kind == "runin":
            p = d.add_paragraph()
            p.add_run(label + " ").bold = True
            p.add_run(q(text))
        else:
            p = d.add_paragraph(q(text))
            if "(" in label or "." in label:  # sub-items are indented
                p.paragraph_format.left_indent = Inches(0.3)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if style != "runin" else WD_ALIGN_PARAGRAPH.LEFT

    d.add_paragraph("IN WITNESS WHEREOF, the parties have signed this Agreement as of the Effective Date."
                    if style != "runin" else "Signed by the parties as of the Effective Date.")
    a, b = (spec.company, spec.counterparty) if spec.company_first else (spec.counterparty, spec.company)
    table = d.add_table(rows=6, cols=2)
    for col, party in enumerate((a, b)):
        cells = [party.full.upper() if style == "dotted" else party.full, "By: ______________________________",
                 "Name:", "Title:", "Date:", f"Notice address: {party.city}"]
        for row, text in enumerate(cells):
            table.cell(row, col).text = text
    path.parent.mkdir(parents=True, exist_ok=True)
    d.save(str(path))
    return refs


def gold_for(spec: NDASpec, topics: dict) -> dict:
    items = _items(spec)
    labels, notes = {}, {}
    for tid, t in topics.items():
        present = {
            1: items["T1"], 2: items["T2"], 3: items["T3"], 4: items["T4"],
            5: items["T5_inline"] if spec.style == "runin" else items["T5"], 6: items["T6"], 7: items["T7"],
            8: items["T8"], 9: items["T9"], 10: items["T10"], 11: items["T11"], 12: items["T12"],
            13: items["T13"], 14: items["T14"],
        }[tid] is not None
        labels[str(tid)] = "acceptable" if present else ("absent" if t.if_absent == "absent" else "escalate")
    for d in spec.deviations:
        labels[str(d.topic)] = d.status
        notes[str(d.topic)] = d.planted
    for tid, note in spec.gold_notes.items():
        notes[str(tid)] = note
    return {"company": spec.company.full, "counterparty": spec.counterparty.full, "labels": labels, "notes": notes}


def main() -> int:
    topics = pb.parse()
    gold_path = DATA / "gold_labels.json"
    if gold_path.exists() and json.loads(gold_path.read_text(encoding="utf8")).get("reviewed"):
        print(f"{gold_path} is marked reviewed; refusing to regenerate the test set over it.")
        return 1

    manifest, gold = [], {}
    for spec in specs():
        refs = write_docx(spec, NDA_DIR / spec.file)
        omitted = [d.topic for d in spec.deviations if d.text is None]
        entry = {
            "file": spec.file, "kind": spec.kind, "style": spec.style,
            "company": spec.company.full, "counterparty": spec.counterparty.full,
            "deviations": [], "omitted_topics": omitted,
        }
        for d in spec.deviations:
            counted = not (d.text is None and d.status == "absent")
            where = "clause removed" if d.text is None else f"section {refs[d.item]}"
            if spec.extra_defs and d.topic == 1:
                where += " (definitions in " + ", ".join(f"section {refs[k]}" for k, _ in spec.extra_defs) + ")"
            entry["deviations"].append({
                "topic_id": d.topic, "topic_name": topics[d.topic].name, "expected_status": d.status,
                "planted": d.planted, "where": where, "counted": counted})
        manifest.append(entry)
        gold[spec.file] = gold_for(spec, topics)
        print(f"{spec.file} {spec.kind:<9} {spec.style:<7} "
              f"{sum(x['counted'] for x in entry['deviations'])} planted")

    (DATA / "seeded_deviations.json").write_text(json.dumps({
        "generated_by": "scripts/make_synthetic.py",
        "note": ("All parties, cities, and projects are fictional. 'company' is the party using the playbook. "
                 "Deviations with counted=false are not planted problems (for example, an omitted clause whose "
                 "absence the playbook accepts)."),
        "ndas": manifest}, indent=2) + "\n", encoding="utf8", newline="\n")

    gold_path.write_text(json.dumps({
        "reviewed": False,
        "playbook_version": pb.version(),
        "labeler": "DRAFT generated from data/seeded_deviations.json; not yet reviewed by a human",
        "instructions": ("The repo owner reads each NDA, corrects any label below, adds a note for each change, and "
                         "then sets reviewed to true. eval.py refuses to run until reviewed is true. Topics with no "
                         "planted deviation default to acceptable, or to the playbook's If absent rule when the "
                         "clause is missing."),
        "ndas": gold}, indent=2) + "\n", encoding="utf8", newline="\n")

    timing = DATA / "manual_timing.csv"
    if not timing.exists():
        with timing.open("w", newline="", encoding="utf8") as f:
            w = csv.writer(f, lineterminator="\n")
            w.writerow(["file", "minutes_manual_review", "minutes_review_of_ai_output"])
            for spec in specs():
                w.writerow([spec.file, "", ""])
    print(f"wrote {len(manifest)} NDAs, seeded_deviations.json, gold_labels.json (reviewed: false), manual_timing.csv")
    return 0


if __name__ == "__main__":
    sys.exit(main())
